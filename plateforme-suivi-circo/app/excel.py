"""Génération du classeur Excel au format des fiches école de la circonscription."""
import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .db import get_db
from .referentiels import CYCLES, NIVEAUX_EFFECTIFS, STATUTS_LIBELLES
from .stats import COLS_EFF, effectifs_par_ecole, personnels_ecole, stats_ecole, synthese_circonscription, total_classe

BLEU = PatternFill("solid", fgColor="1F4E79")
BLEU_CLAIR = PatternFill("solid", fgColor="DDEBF7")
GRIS = PatternFill("solid", fgColor="F2F2F2")
VERT = PatternFill("solid", fgColor="E2EFDA")
JAUNE = PatternFill("solid", fgColor="FFF2CC")
BLANC_GRAS = Font(bold=True, color="FFFFFF")
GRAS = Font(bold=True)
TITRE = Font(bold=True, size=14, color="1F4E79")
FIN = Side(style="thin", color="BFBFBF")
BORDURE = Border(left=FIN, right=FIN, top=FIN, bottom=FIN)
CENTRE = Alignment(horizontal="center", vertical="center", wrap_text=True)
GAUCHE = Alignment(horizontal="left", vertical="center", wrap_text=True)
LIBELLES_NIVEAUX = [lib for _, lib, _ in NIVEAUX_EFFECTIFS]


def _titre(ws, ligne, texte, nb_col, fill=BLEU, font=BLANC_GRAS):
    ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=nb_col)
    c = ws.cell(row=ligne, column=1, value=texte)
    c.fill, c.font, c.alignment = fill, font, GAUCHE
    ws.row_dimensions[ligne].height = 22


def _entetes(ws, ligne, valeurs, col=1, fill=BLEU_CLAIR):
    for i, v in enumerate(valeurs):
        c = ws.cell(row=ligne, column=col + i, value=v)
        c.fill, c.font, c.alignment, c.border = fill, GRAS, CENTRE, BORDURE


def _ligne(ws, ligne, valeurs, col=1, gras=False, fill=None):
    for i, v in enumerate(valeurs):
        c = ws.cell(row=ligne, column=col + i, value=v)
        c.border = BORDURE
        c.alignment = GAUCHE if isinstance(v, str) and i == 0 else CENTRE
        if gras:
            c.font = GRAS
        if fill:
            c.fill = fill


def _champ(ws, ligne, libelle, valeur, col=1, largeur=3):
    c = ws.cell(row=ligne, column=col, value=libelle)
    c.font, c.fill, c.border, c.alignment = GRAS, GRIS, BORDURE, GAUCHE
    ws.merge_cells(start_row=ligne, start_column=col + 1, end_row=ligne, end_column=col + largeur)
    v = ws.cell(row=ligne, column=col + 1, value=valeur)
    v.border, v.alignment = BORDURE, GAUCHE
    for k in range(col + 2, col + largeur + 1):
        ws.cell(row=ligne, column=k).border = BORDURE


def _nom_feuille(nom, existants):
    base = re.sub(r"[\[\]\*\?/\\:]", " ", nom).strip()[:28] or "Ecole"
    candidat, n = base, 2
    while candidat.upper() in existants:
        candidat = f"{base[:25]} {n}"
        n += 1
    existants.add(candidat.upper())
    return candidat


def _feuille_bord(wb, params, s):
    ws = wb.active
    ws.title = "BORD"
    ws.column_dimensions["A"].width = 26
    for col in "BCDEFGHIJ":
        ws.column_dimensions[col].width = 14
    _titre(ws, 1, f"{params.get('academie', '')}  —  {params.get('circonscription', 'CIRCONSCRIPTION')}", 10)
    _titre(ws, 2, f"TABLEAU DE BORD  —  Année scolaire {params.get('annee_scolaire', '')}", 10, BLEU_CLAIR, TITRE)

    _entetes(ws, 4, ["ÉCOLES", "ÉLÈVES (hors ULIS)", "ÉLÈVES ULIS", "CLASSES", "ENSEIGNANTS", "PERSONNELS"])
    _ligne(ws, 5, [s["nb_ecoles"], s["eleves_hors_ulis"], s["eleves_ulis"], s["classes"], s["enseignants"],
                   s["personnels"]], gras=True)

    _titre(ws, 7, "PAR CYCLE", 6, BLEU_CLAIR, GRAS)
    _entetes(ws, 8, [CYCLES[1], CYCLES[2], CYCLES[3], "ULIS", "TOTAL"])
    _ligne(ws, 9, [s["cycles"][1], s["cycles"][2], s["cycles"][3], s["eleves_ulis"], s["eleves_hors_ulis"]], gras=True)

    _titre(ws, 11, "PAR NIVEAU", 6, BLEU_CLAIR, GRAS)
    _entetes(ws, 12, ["Niveau", "Élèves", "Classes", "Élèves / classe"])
    r = 13
    for n in s["par_niveau"]:
        _ligne(ws, r, [n["niveau"], n["eleves"], n["classes"], n["ratio"] if n["ratio"] is not None else "—"])
        r += 1
    _ligne(ws, r, ["TOTAL", s["eleves_hors_ulis"], s["classes"],
                   round(s["eleves_hors_ulis"] / s["classes"], 1) if s["classes"] else "—"], gras=True, fill=VERT)


def _feuille_effectifs(wb, params):
    ws = wb.create_sheet("EFFECTIFS PAR NIVEAUX")
    ecoles, totaux = effectifs_par_ecole()
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 12
    for i in range(3, 17):
        ws.column_dimensions[get_column_letter(i)].width = 9
    _titre(ws, 1, "EFFECTIFS PAR ÉCOLE ET PAR NIVEAU", 16)
    _entetes(ws, 3, ["ÉCOLES", "TYPE", "PS", "MS", "GS", "CP", "CE1", "CE2", "CM1", "CM2", "ULIS",
                     "CYCLE 1", "CYCLE 2", "CYCLE 3", "TOTAL", "CLASSES"])
    r = 4
    for e in ecoles:
        _ligne(ws, r, [e["nom"], e["type"], *[e[c] or "" for c in COLS_EFF], e["cycle1"], e["cycle2"], e["cycle3"],
                       e["total"], e["nb_classes_calc"]])
        r += 1
    _ligne(ws, r, ["TOTAL CIRCONSCRIPTION", "", *[totaux[c] for c in COLS_EFF], totaux["cycle1"], totaux["cycle2"],
                   totaux["cycle3"], totaux["total"], totaux["nb_classes_calc"]], gras=True, fill=VERT)


def _feuille_ecole(wb, params, ecole, existants):
    ws = wb.create_sheet(_nom_feuille(ecole["nom"], existants))
    dans_org, hors_org = personnels_ecole(ecole["id"])
    st = stats_ecole(ecole, dans_org, hors_org)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 11
    ws.column_dimensions["E"].width = 8
    for i in range(6, 16):
        ws.column_dimensions[get_column_letter(i)].width = 7.5

    _titre(ws, 1, f"{params.get('academie', '')}  —  {params.get('circonscription', '')}", 15)
    _titre(ws, 2, f"{ecole['nom'].upper()}   ({ecole['type']})   —   Année scolaire {params.get('annee_scolaire', '')}",
           15, BLEU_CLAIR, TITRE)

    r = 4
    _titre(ws, r, "📋  FICHE SIGNALÉTIQUE", 15, GRIS, GRAS)
    champs = [
        ("RNE :", ecole["rne"]), ("Adresse :", ecole["adresse"]),
        ("Téléphone / Courriel :", " / ".join(v for v in (ecole["telephone"], ecole["email"]) if v)),
        ("Horaires :", ecole["horaires"]), ("APC & horaires :", ecole["apc"]),
        ("Directeur / Directrice :", ecole["directeur"]), ("Secrétaire :", ecole["secretaire"]),
        ("Organisation (OTS) :", ecole["ots"]),
        ("FAEX :", " / ".join(v for v in (ecole["faex_nom"], ecole["faex_telephone"], ecole["faex_email"]) if v)),
        ("% intervention FAEX :", ecole["faex_pct"]),
    ]
    aesh = [p for p in hors_org if p["statut"] == "AESH"]
    for i, p in enumerate(aesh, 1):
        champs.append((f"AESH {i} :", " / ".join(v for v in (f"{p['nom']} {p['prenom']}".strip(), p["telephone"]) if v)))
    for lib, val in champs:
        r += 1
        _champ(ws, r, lib, val, largeur=14)

    r += 2
    _titre(ws, r, "🏫  INFORMATIONS GÉNÉRALES", 15, GRIS, GRAS)
    r += 1
    _champ(ws, r, "Nombre de salles :", ecole["nb_salles"], largeur=3)
    _champ(ws, r, "Nombre total de classes :", ecole["nb_classes"], col=6, largeur=3)
    _champ(ws, r, "Éducation prioritaire :", ecole["education_prioritaire"], col=11, largeur=4)
    r += 1
    _champ(ws, r, "% décharge de direction :", ecole["decharge_direction"], largeur=3)
    _champ(ws, r, "Date de mise à jour :", ecole["date_maj"], col=6, largeur=3)

    r += 2
    _titre(ws, r, "👩‍🏫  ORGANISATION PÉDAGOGIQUE", 15, GRIS, GRAS)
    r += 1
    _entetes(ws, r, ["Enseignant(e)", "Statut", "Sexe", "Dispositif", "Salle", *LIBELLES_NIVEAUX, "TOTAL*"])
    debut = r + 1
    for p in dans_org:
        r += 1
        _ligne(ws, r, [f"{p['nom']} {p['prenom']}".strip(), p["statut"], p["sexe"], p["dispositif"], p["salle"],
                       *[p[c] or "" for c in COLS_EFF], total_classe(p)])
    r += 1
    _ligne(ws, r, ["EFFECTIFS TOTAUX PAR NIVEAU", "", "", "", "", *[st["effectifs"][c] for c in COLS_EFF],
                   st["total_eleves"]], gras=True, fill=VERT)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)

    r += 2
    _titre(ws, r, "👨‍🏫  PERSONNELS NON COMPTABILISÉS DANS L'ORGANISATION PÉDAGOGIQUE (TR-ZIL / BGS / AESH…)",
           15, GRIS, GRAS)
    r += 1
    _entetes(ws, r, ["Nom et prénom", "Statut", "Sexe", "Anc. (ans)", "Fonction", "Observations"])
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=15)
    for p in hors_org:
        r += 1
        _ligne(ws, r, [f"{p['nom']} {p['prenom']}".strip(), p["statut"], p["sexe"],
                       p["anciennete"] if p["anciennete"] is not None else "", p["fonction"], p["observations"]])
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=15)

    r += 2
    _titre(ws, r, "📊  STATISTIQUES — calculées automatiquement", 15, GRIS, GRAS)
    r += 1
    _entetes(ws, r, ["Nbre de classes par niveau", "", "", "", "", *LIBELLES_NIVEAUX, "TOTAL"])
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    r += 1
    _ligne(ws, r, ["", "", "", "", "", *[st["classes"][c] for c in COLS_EFF], st["total_classes"]], gras=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    r += 1
    _champ(ws, r, "Nbre total d'enseignants :", st["nb_enseignants"], largeur=3)
    _champ(ws, r, "Ancienneté moyenne (ans) :", st["anciennete_moyenne"] if st["anciennete_moyenne"] is not None else "—",
           col=6, largeur=3)
    _champ(ws, r, "Élèves / classe :", st["eleves_par_classe"] if st["eleves_par_classe"] is not None else "—",
           col=11, largeur=4)

    r += 2
    _entetes(ws, r, ["Tranche d'ancienneté", *[lib for lib, _ in st["tranches"]], "TOTAL"])
    r += 1
    _ligne(ws, r, ["Nb enseignants", *[n for _, n in st["tranches"]], sum(n for _, n in st["tranches"])])

    r += 2
    _entetes(ws, r, ["Statut", "Hommes", "Femmes", "Total", "Libellé"])
    ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=10)
    for s in st["statuts"]:
        r += 1
        _ligne(ws, r, [s["statut"], s["hommes"], s["femmes"], s["total"], STATUTS_LIBELLES.get(s["statut"], "")])
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=10)
    r += 1
    _ligne(ws, r, ["TOTAL", st["total_hommes"], st["total_femmes"], st["nb_enseignants"], ""], gras=True, fill=VERT)

    r += 2
    _entetes(ws, r, ["Dispositif", "Nb enseignants", "Nb cl. CP", "Nb cl. CE1"])
    for d in st["dispositifs"]:
        r += 1
        _ligne(ws, r, [d["dispositif"], d["nb"], d["cp"] if d["dispositif"] == "Duo" else "—",
                       d["ce1"] if d["dispositif"] == "Duo" else "—"])

    r += 2
    _titre(ws, r, "✅  Contrôles de cohérence automatiques", 15, GRIS, GRAS)
    r += 1
    _entetes(ws, r, ["Contrôle", "Résultat", "Valeur A", "Valeur B"])
    for c in st["controles"]:
        r += 1
        _ligne(ws, r, [c["libelle"], "✔ OK" if c["ok"] else "⚠ Écart", c["a"], c["b"]],
               fill=None if c["ok"] else JAUNE)
    r += 2
    ws.cell(row=r, column=1, value="(*) TOTAL : effectifs hors ULIS (élèves ULIS déjà comptabilisés dans leur classe "
                                    "ordinaire). APC : Activités Pédagogiques Complémentaires. OTS : Organisation du "
                                    "Temps Scolaire.").font = Font(italic=True, size=9)
    ws.freeze_panes = ws.cell(row=debut, column=2)


def construire_classeur(params):
    """Retourne un tampon BytesIO contenant le classeur complet."""
    wb = Workbook()
    synthese = synthese_circonscription()
    _feuille_bord(wb, params, synthese)
    _feuille_effectifs(wb, params)
    existants = {"BORD", "EFFECTIFS PAR NIVEAUX"}
    for ecole in get_db().execute("SELECT * FROM ecoles ORDER BY commune, nom").fetchall():
        _feuille_ecole(wb, params, ecole, existants)
    tampon = io.BytesIO()
    wb.save(tampon)
    tampon.seek(0)
    return tampon


# ---------------------------------------------------------------- carte scolaire ---

def _feuille_synthese_carte(wb, params, campagne, groupes, general):
    ws = wb.active
    ws.title = "Synthèse circo"
    ws.column_dimensions["A"].width = 34
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 17
    ac = params.get("campagne_constat") or "Constat"
    _titre(ws, 1, f"CARTE SCOLAIRE {campagne} — SYNTHÈSE DE CIRCONSCRIPTION", 7)
    _titre(ws, 2, f"{params.get('circonscription', '')} — IEN : {params.get('ien', '')}", 7, BLEU_CLAIR, GRAS)
    ws.cell(row=3, column=1,
            value=f"Constat de rentrée → prévisions {campagne} (avant / après mesures de carte scolaire)").font = Font(italic=True, size=9)

    r = 5
    _titre(ws, r, "EFFECTIFS", 7, GRIS, GRAS)
    r += 1
    _entetes(ws, r, ["", f"{ac}", f"Prévisions {campagne}", "Évolution", "Évolution %"])
    for g in groupes:
        r += 1
        t = g["totaux"]
        _ligne(ws, r, [g["titre"], t["total_c"], t["total_p"], t["evolution"],
                       (t["pct"] / 100) if t["pct"] is not None else ""])
        ws.cell(row=r, column=5).number_format = "0.0%"
    r += 1
    _ligne(ws, r, ["TOTAL CIRCONSCRIPTION", general["total_c"], general["total_p"], general["evolution"],
                   (general["pct"] / 100) if general["pct"] is not None else ""], gras=True, fill=VERT)
    ws.cell(row=r, column=5).number_format = "0.0%"

    r += 2
    _titre(ws, r, "DIVISIONS (CLASSES)", 7, GRIS, GRAS)
    r += 1
    _entetes(ws, r, ["", "Constat", "Après mesures", "Ouvertures", "Fermetures", "Solde"])
    for g in groupes:
        r += 1
        t = g["totaux"]
        _ligne(ws, r, [g["titre"], t["div"], t["div_apres"], t["ouverture"], t["fermeture"],
                       t["ouverture"] - t["fermeture"]])
    r += 1
    _ligne(ws, r, ["TOTAL CIRCONSCRIPTION", general["div"], general["div_apres"], general["ouverture"],
                   general["fermeture"], general["ouverture"] - general["fermeture"]], gras=True, fill=VERT)

    r += 2
    _titre(ws, r, "E/D MOYEN (effectifs / divisions)", 7, GRIS, GRAS)
    r += 1
    _entetes(ws, r, ["", "E/D constat", "E/D projeté avant mesures", "E/D après mesures"])
    for g in groupes:
        r += 1
        t = g["totaux"]
        _ligne(ws, r, [g["titre"], t["ed"] or "—", t["ed_avant"] or "—", t["ed_apres"] or "—"])
    r += 1
    _ligne(ws, r, ["TOTAL CIRCONSCRIPTION", general["ed"] or "—", general["ed_avant"] or "—",
                   general["ed_apres"] or "—"], gras=True, fill=VERT)

    r += 2
    _titre(ws, r, "SALLES DE CLASSE", 7, GRIS, GRAS)
    r += 1
    _entetes(ws, r, ["", "Salles disponibles", f"Salles nécessaires {campagne}", "Écart"])
    for g in groupes:
        r += 1
        t = g["totaux"]
        _ligne(ws, r, [g["titre"], t["salles"], t["salles_necessaires"], t["ecart_salles"]],
               fill=None if t["ecart_salles"] >= 0 else JAUNE)
    r += 1
    _ligne(ws, r, ["TOTAL CIRCONSCRIPTION", general["salles"], general["salles_necessaires"],
                   general["ecart_salles"]], gras=True, fill=VERT)

    r += 2
    for texte in [
        "Mode d'emploi",
        "1. Le constat de rentrée est repris des fiches école (organisation pédagogique). Il peut être saisi à la main école par école.",
        "2. Les prévisions se calculent par montée pédagogique (MS = PS, GS = MS, CE1 = CP…).",
        "   Seuls les PS prévus (maternelles et primaires) et les CP prévus (élémentaires, selon les GS des maternelles de rattachement) sont saisis.",
        "3. Les ouvertures / fermetures proposées recalculent les divisions après mesures, les E/D et les besoins en salles.",
        "4. Cette synthèse peut être transmise en l'état pour le dialogue de carte scolaire.",
    ]:
        ws.cell(row=r, column=1, value=texte).font = Font(italic=(texte != "Mode d'emploi"), size=9,
                                                          bold=(texte == "Mode d'emploi"))
        r += 1


def _feuille_type_carte(wb, params, campagne, groupe):
    from .carte import niveaux_du_type
    ws = wb.create_sheet(groupe["type"][:28])
    lignes, t = groupe["lignes"], groupe["totaux"]
    niveaux = niveaux_du_type(groupe["type"])
    libelles = {n: lib for n, lib in [("ps", "PS"), ("ms", "MS"), ("gs", "GS"), ("cp", "CP"),
                                      ("ce1", "CE1"), ("ce2", "CE2"), ("cm1", "CM1"), ("cm2", "CM2")]}
    sup = (["Div. CP-CE1", "Div. CE2-CM1-CM2", "dont rotation / duo"] if groupe["type"] == "Élémentaire"
           else ["Div. mat.", "Div. élém."] if groupe["type"] == "Primaire" else [])
    cles_sup = (["div_cp_ce1", "div_ce2_cm1_cm2", "div_rotation"] if groupe["type"] == "Élémentaire"
                else ["div_mat", "div_elem"] if groupe["type"] == "Primaire" else [])
    entetes = (["École", "RNE"] + [libelles[n] for n in niveaux] + ["Effectif total", "Divisions"] + sup
               + [f"{libelles[n]} prévus" for n in niveaux]
               + ["Effectif total prévu", "E/D projeté avant mesures", "Ouverture", "Fermeture",
                  "Divisions après mesures", "E/D après mesures", "Évolution effectifs",
                  "Nb de salles", "Salles nécessaires", "Écart salles", "Observations"])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 12
    for i in range(3, len(entetes) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 11
    ws.column_dimensions[get_column_letter(len(entetes))].width = 30

    _titre(ws, 1, f"CARTE SCOLAIRE {campagne} — {groupe['titre'].upper()}", len(entetes))
    _titre(ws, 2, f"{params.get('circonscription', '')} — constat de rentrée → prévisions {campagne}",
           len(entetes), BLEU_CLAIR, GRAS)
    _entetes(ws, 4, entetes)
    r = 4
    for l in lignes:
        r += 1
        _ligne(ws, r, [l["ecole"]["nom"], l["ecole"]["rne"]]
               + [l["constat"][n] for n in niveaux] + [l["total_c"], l["div"]]
               + [l[c] for c in cles_sup]
               + [l["prev"][n] for n in niveaux]
               + [l["total_p"], l["ed_avant"] or "—", l["ouverture"] or "", l["fermeture"] or "",
                  l["div_apres"], l["ed_apres"] or "—", l["evolution"], l["salles"],
                  l["salles_necessaires"], l["ecart_salles"], l["carte"]["observations"]])
    r += 1
    _ligne(ws, r, [f"TOTAL {groupe['titre'].upper()}", ""]
           + [t["constat"][n] for n in niveaux] + [t["total_c"], t["div"]]
           + [t[c] for c in cles_sup]
           + [t["prev"][n] for n in niveaux]
           + [t["total_p"], t["ed_avant"] or "—", t["ouverture"], t["fermeture"], t["div_apres"],
              t["ed_apres"] or "—", t["evolution"], t["salles"], t["salles_necessaires"],
              t["ecart_salles"], ""], gras=True, fill=VERT)
    ws.freeze_panes = ws.cell(row=5, column=3)


def _feuille_postes(wb, params, campagne, lignes, totaux_p):
    ws = wb.create_sheet("Postes hors classe")
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 18
    for col in "CDEFG":
        ws.column_dimensions[col].width = 16
    _titre(ws, 1, "POSTES DE PROFESSEURS — HORS DE LA CLASSE", 7)
    _titre(ws, 2, f"{params.get('circonscription', '')} — campagne carte scolaire {campagne}", 7, BLEU_CLAIR, GRAS)
    _entetes(ws, 4, ["Poste", "Spécialité", "Supports (constat)", "Affectations (constat)",
                     "Ouverture", "Fermeture", f"Supports {campagne}"])
    r = 4
    for l in lignes:
        r += 1
        _ligne(ws, r, [l["poste"], l["specialite"], l["supports"], l["affectations"],
                       l["ouverture"] or "", l["fermeture"] or "", l["apres"]])
    r += 1
    _ligne(ws, r, ["TOTAL", "", totaux_p["supports"], totaux_p["affectations"], totaux_p["ouverture"],
                   totaux_p["fermeture"], totaux_p["apres"]], gras=True, fill=VERT)


def construire_carte_scolaire(params, campagne):
    """Classeur de carte scolaire : synthèse, une feuille par type d'école, postes hors classe."""
    from .carte import lignes_campagne, postes, totaux, totaux_postes
    groupes = []
    for type_ecole, titre in [("Maternelle", "Écoles maternelles"), ("Élémentaire", "Écoles élémentaires"),
                              ("Primaire", "Écoles primaires")]:
        lignes = lignes_campagne(campagne, type_ecole)
        groupes.append({"type": type_ecole, "titre": titre, "lignes": lignes, "totaux": totaux(lignes)})
    general = totaux([l for g in groupes for l in g["lignes"]])

    wb = Workbook()
    _feuille_synthese_carte(wb, params, campagne, groupes, general)
    for g in groupes:
        _feuille_type_carte(wb, params, campagne, g)
    lignes_postes = postes(campagne)
    _feuille_postes(wb, params, campagne, lignes_postes, totaux_postes(lignes_postes))
    tampon = io.BytesIO()
    wb.save(tampon)
    tampon.seek(0)
    return tampon
