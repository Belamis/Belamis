"""Import d'un classeur Excel au format « fiche école » de la circonscription.

Le classeur attendu contient une feuille par école avec, en colonne A, les libellés
« RNE : », « Adresse… », « Directeur : »… puis un tableau « Organisation pédagogique »
(Enseignant(e) / Statut / Sexe / Dispositif / Salle / PS … CM2 / ULIS / TOTAL) et un
tableau des personnels non comptabilisés (TR-ZIL, BGS, AESH…).
"""
import datetime as dt
import re
import unicodedata

from openpyxl import load_workbook

from .db import get_db
from .referentiels import STATUTS_CODES
from .stats import COLS_EFF

CIVILITES = re.compile(r"^(mme|mr|m\.|m|mlle|melle|monsieur|madame)\s+", re.I)


def _norm(texte):
    """Minuscules sans accents, espaces réduits — pour comparer des libellés."""
    if texte is None:
        return ""
    t = unicodedata.normalize("NFKD", str(texte)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", t).strip().lower()


def _texte(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def _entier(v):
    if v is None or v == "":
        return 0
    try:
        return int(float(str(v).replace(",", ".")))
    except ValueError:
        return 0


def _pourcentage(v):
    if v is None or v == "":
        return 0
    try:
        f = float(str(v).replace("%", "").replace(",", "."))
    except ValueError:
        return 0
    return int(round(f * 100)) if f <= 1 else int(round(f))


def _premier(ws, ligne, colonnes):
    for col in colonnes:
        v = ws.cell(row=ligne, column=col).value
        if v not in (None, "", " ") and not _texte(v).startswith(("→", "->")):
            return v
    return None


def separer_nom_prenom(complet):
    """« BOINALI Habachia » -> ('BOINALI', 'Habachia') ; « Mme SAID ALI Rafia » -> ('SAID ALI', 'Rafia')."""
    complet = CIVILITES.sub("", _texte(complet)).strip()
    complet = re.sub(r"\s*\(.*?\)\s*$", "", complet)  # « (C7) » en fin de nom
    mots = complet.split()
    if not mots:
        return "", ""
    nom = []
    while mots and (mots[0].isupper() or (len(mots[0]) > 1 and mots[0].replace("'", "").replace("’", "").isupper())):
        nom.append(mots.pop(0))
    if not nom:  # tout en minuscules / capitalisé : premier mot = nom
        nom.append(mots.pop(0))
    return " ".join(nom).upper(), " ".join(mots)


def _statut(brut, defaut="AUTRE"):
    s = _texte(brut).upper().replace(" ", "")
    if s in STATUTS_CODES:
        return s, ""
    correspondances = {"S": "STGF", "STAGIAIRE": "STGF", "TRZIL": "TR-ZIL", "TR": "TR-ZIL", "ZIL": "TR-ZIL",
                       "BRIGADE": "BGS", "RGS": "BGS", "CDI": "C/CDI", "CONTRACTUEL": "C", "TITULAIRE": "PE"}
    if s in correspondances:
        return correspondances[s], ""
    if "AESH" in s:
        return "AESH", _texte(brut)
    return defaut, _texte(brut)


def _sexe(brut):
    s = _texte(brut).upper()[:1]
    return {"F": "F", "H": "H", "M": "H"}.get(s, "")


def _type_ecole(ws, nom):
    for col in range(2, 16):
        if _texte(ws.cell(row=3, column=col).value).upper() == "X":
            lib = _norm(ws.cell(row=3, column=col + 1).value)
            if "mater" in lib:
                return "Maternelle"
            if "elem" in lib:
                return "Élémentaire"
            if "prim" in lib:
                return "Primaire"
    n = _norm(nom)
    if "mat" in n:
        return "Maternelle"
    if "elem" in n:
        return "Élémentaire"
    return "Primaire"


def est_fiche_ecole(ws):
    for r in range(1, 12):
        if "fiche signaletique" in _norm(ws.cell(row=r, column=1).value):
            return True
    return False


def lire_fiche(ws):
    """Extrait une fiche école (dict) + personnels (liste de dicts) d'une feuille."""
    ecole = {"nom": _texte(ws["A2"].value) or ws.title, "type": _type_ecole(ws, ws.title)}
    annee = ""
    for col in range(8, 16):
        v = _texte(ws.cell(row=2, column=col).value)
        if re.match(r"^\d{4}\s*[/-]\s*\d{4}$", v):
            annee = v
    entete = _texte(ws["A1"].value)
    aesh, personnels = [], []

    # --- Signalétique : repérage par libellé en colonne A --------------------
    for r in range(4, 30):
        lib = _norm(ws.cell(row=r, column=1).value)
        val = _premier(ws, r, range(2, 8))
        if not lib:
            continue
        if lib.startswith("rne"):
            ecole["rne"] = _texte(val).upper()
        elif lib.startswith("adresse"):
            ecole["adresse"] = _texte(val)
        elif lib.startswith("telephone / courriel") or lib.startswith("telephone"):
            morceaux = [m.strip() for m in _texte(val).split("/")]
            ecole["email"] = " / ".join(m for m in morceaux if "@" in m)
            ecole["telephone"] = " / ".join(m for m in morceaux if "@" not in m and m)
        elif lib.startswith("horaires"):
            ecole["horaires"] = _texte(val)
        elif lib.startswith("apc"):
            ecole["apc"] = _texte(val)
        elif lib.startswith("directeur") or lib.startswith("directrice"):
            ecole["directeur"] = _texte(val)
        elif lib.startswith("secretaire"):
            ecole["secretaire"] = _texte(val)
        elif lib.startswith("organisation"):
            ecole["ots"] = _texte(val)
        elif lib.startswith("faex"):
            ecole["faex_nom"] = _texte(val)
        elif lib.startswith("tel. faex") or lib.startswith("tel faex"):
            ecole["faex_telephone"] = _texte(val)
            for col in range(8, 16):
                if "courriel" in _norm(ws.cell(row=r, column=col).value):
                    ecole["faex_email"] = _texte(_premier(ws, r, range(col + 1, col + 6)))
        elif "intervention faex" in lib:
            ecole["faex_pct"] = _texte(val)
        elif lib.startswith("aesh"):
            nom_complet = _texte(val)
            tel = ""
            for col in range(8, 16):
                if "tel" in _norm(ws.cell(row=r, column=col).value):
                    tel = _texte(_premier(ws, r, range(col + 1, col + 6)))
            if nom_complet and not nom_complet.isdigit():
                nom, prenom = separer_nom_prenom(nom_complet)
                aesh.append({"nom": nom, "prenom": prenom, "statut": "AESH", "fonction": "AESH", "sexe": "",
                             "telephone": tel, "dans_organisation": 0, "anciennete": None, "observations": ""})
        elif lib.startswith("nombre de salles"):
            ecole["nb_salles"] = _entier(re.sub(r"[^\d].*$", "", _texte(val)) or 0)
            for col in range(6, 16):
                if "classes" in _norm(ws.cell(row=r, column=col).value):
                    droite = _premier(ws, r, range(col + 1, col + 6))
                    gauche = ws.cell(row=r, column=col - 1).value
                    ecole["nb_classes"] = _entier(droite if droite is not None else gauche)
        elif lib.startswith("% decharge"):
            ecole["decharge_direction"] = _pourcentage(val)
            for col in range(6, 16):
                if "mise a jour" in _norm(ws.cell(row=r, column=col).value):
                    ecole["date_maj"] = _texte(_premier(ws, r, range(col + 1, col + 6)))[:10]
        elif lib.startswith("enseignant"):
            break

    # --- Organisation pédagogique ---------------------------------------------
    ligne_entete = None
    for r in range(4, ws.max_row + 1):
        if _norm(ws.cell(row=r, column=1).value).startswith("enseignant(e)"):
            ligne_entete = r
            break
    if ligne_entete:
        colonnes = {}
        for col in range(1, 20):
            lib = _norm(ws.cell(row=ligne_entete, column=col).value)
            if lib:
                colonnes[lib] = col
        col_niveau = {c: colonnes.get(lib.lower()) for c, lib, _ in
                      [("eff_ps", "ps", 1), ("eff_ms", "ms", 1), ("eff_gs", "gs", 1), ("eff_cp", "cp", 2),
                       ("eff_ce1", "ce1", 2), ("eff_ce2", "ce2", 2), ("eff_cm1", "cm1", 3), ("eff_cm2", "cm2", 3),
                       ("eff_ulis", "ulis", 0)]}
        for r in range(ligne_entete + 1, ws.max_row + 1):
            a = _texte(ws.cell(row=r, column=1).value)
            if _norm(a).startswith("effectifs totaux") or _norm(a).startswith("statistiques"):
                break
            if not a:
                continue
            nom, prenom = separer_nom_prenom(a)
            statut, fonction = _statut(ws.cell(row=r, column=colonnes.get("statut", 2)).value, defaut="PE")
            p = {"nom": nom, "prenom": prenom, "statut": statut, "fonction": fonction,
                 "sexe": _sexe(ws.cell(row=r, column=colonnes.get("sexe", 3)).value),
                 "dispositif": _texte(ws.cell(row=r, column=colonnes.get("dispositif", 4)).value).capitalize(),
                 "salle": _texte(ws.cell(row=r, column=colonnes.get("salle", 5)).value),
                 "dans_organisation": 1, "anciennete": None, "telephone": "", "observations": ""}
            if p["dispositif"] not in ("Solo", "Duo"):
                p["dispositif"] = ""
            for c, col in col_niveau.items():
                p[c] = _entier(ws.cell(row=r, column=col).value) if col else 0
            personnels.append(p)

    # --- Personnels non comptabilisés ------------------------------------------
    for r in range(4, ws.max_row + 1):
        if "non comptabilis" in _norm(ws.cell(row=r, column=1).value):
            ligne_hors = r + 1
            cols = {}
            for col in range(1, 20):
                lib = _norm(ws.cell(row=ligne_hors, column=col).value)
                if lib:
                    cols[lib] = col
            col_nom = cols.get("nom et prenom", 2)
            col_sexe = cols.get("sexe", 7)
            col_anc = next((c for l, c in cols.items() if l.startswith("anc")), 9)
            col_fonction = cols.get("fonction", 11)
            col_obs = cols.get("observations", 13)
            for rr in range(ligne_hors + 1, ws.max_row + 1):
                if _norm(ws.cell(row=rr, column=1).value).startswith("statistiques"):
                    break
                nom_complet = _texte(ws.cell(row=rr, column=col_nom).value)
                if not nom_complet:
                    continue
                nom, prenom = separer_nom_prenom(nom_complet)
                statut, fonction = _statut(ws.cell(row=rr, column=col_fonction).value)
                anc = _texte(ws.cell(row=rr, column=col_anc).value)
                personnels.append({
                    "nom": nom, "prenom": prenom, "statut": statut,
                    "fonction": fonction or _texte(ws.cell(row=rr, column=col_fonction).value),
                    "sexe": _sexe(ws.cell(row=rr, column=col_sexe).value),
                    "anciennete": int(anc) if anc.isdigit() else None, "dans_organisation": 0,
                    "observations": _texte(ws.cell(row=rr, column=col_obs).value), "telephone": "",
                })
            break

    # Les AESH de la signalétique s'ajoutent s'ils ne figurent pas déjà dans la liste hors organisation.
    deja = {(p["nom"], _norm(p["prenom"])) for p in personnels}
    for a in aesh:
        if (a["nom"], _norm(a["prenom"])) not in deja:
            personnels.append(a)
    return ecole, personnels, annee, entete


def importer_classeur(chemin_ou_flux):
    """Importe toutes les fiches école du classeur. Retourne un rapport (dict)."""
    wb = load_workbook(chemin_ou_flux, data_only=True, read_only=False)
    db = get_db()
    rapport = {"ecoles_creees": 0, "ecoles_maj": 0, "personnels_crees": 0, "personnels_maj": 0,
               "feuilles_ignorees": [], "details": []}
    cols_perso = ["nom", "prenom", "sexe", "telephone", "statut", "fonction", "anciennete", "dans_organisation",
                  "dispositif", "salle", "observations", *COLS_EFF]

    for ws in wb.worksheets:
        if not est_fiche_ecole(ws):
            rapport["feuilles_ignorees"].append(ws.title)
            continue
        ecole, personnels, annee, entete = lire_fiche(ws)
        if annee:
            db.execute("INSERT INTO parametres (cle, valeur) VALUES ('annee_scolaire', ?) "
                       "ON CONFLICT(cle) DO UPDATE SET valeur = excluded.valeur", (annee.replace(" ", ""),))
        if entete:
            morceaux = [m.strip() for m in re.split(r"[—–-]{1,2}", entete) if m.strip()]
            if len(morceaux) >= 2:
                for cle, val in (("academie", morceaux[0]), ("circonscription", morceaux[1])):
                    db.execute("INSERT INTO parametres (cle, valeur) VALUES (?, ?) "
                               "ON CONFLICT(cle) DO UPDATE SET valeur = excluded.valeur", (cle, val.title()))

        existante = None
        if ecole.get("rne"):
            existante = db.execute("SELECT id FROM ecoles WHERE rne = ?", (ecole["rne"],)).fetchone()
        if existante is None:
            existante = db.execute("SELECT id FROM ecoles WHERE lower(nom) = lower(?)", (ecole["nom"],)).fetchone()
        champs = {k: v for k, v in ecole.items()}
        if existante:
            ecole_id = existante["id"]
            db.execute("UPDATE ecoles SET " + ", ".join(f"{k} = :{k}" for k in champs)
                       + ", modifie_le = datetime('now') WHERE id = :id", {**champs, "id": ecole_id})
            rapport["ecoles_maj"] += 1
        else:
            cur = db.execute(f"INSERT INTO ecoles ({', '.join(champs)}) VALUES ({', '.join(':' + k for k in champs)})",
                             champs)
            ecole_id = cur.lastrowid
            rapport["ecoles_creees"] += 1

        crees = maj = 0
        for p in personnels:
            p.setdefault("dispositif", "")
            p.setdefault("salle", "")
            for c in COLS_EFF:
                p.setdefault(c, 0)
            p["ecole_id"] = ecole_id
            trouve = db.execute(
                "SELECT id FROM enseignants WHERE ecole_id = ? AND upper(nom) = upper(?) AND lower(prenom) = lower(?)",
                (ecole_id, p["nom"], p["prenom"])).fetchone()
            if trouve:
                p["id"] = trouve["id"]
                db.execute("UPDATE enseignants SET " + ", ".join(f"{c} = :{c}" for c in cols_perso)
                           + ", actif = 1, modifie_le = datetime('now') WHERE id = :id", p)
                maj += 1
            else:
                db.execute(f"INSERT INTO enseignants (ecole_id, {', '.join(cols_perso)}) "
                           f"VALUES (:ecole_id, {', '.join(':' + c for c in cols_perso)})", p)
                crees += 1
        rapport["personnels_crees"] += crees
        rapport["personnels_maj"] += maj
        rapport["details"].append({"feuille": ws.title, "ecole": ecole["nom"], "type": ecole["type"],
                                   "personnels": len(personnels), "crees": crees, "maj": maj})
    db.commit()
    return rapport
