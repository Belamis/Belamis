"""Tests de bout en bout : authentification, saisie, statistiques, exports, import Excel."""
import io

from openpyxl import Workbook, load_workbook

from app.importer import separer_nom_prenom
from tests.conftest import Session


def test_login_requis(client):
    r = client.get("/")
    assert r.status_code == 302 and "/login" in r.headers["Location"]


def test_login_mauvais_mot_de_passe(client):
    s = Session(client)
    r = s.login("admin", "faux")
    assert r.status_code == 200 and "incorrect" in r.get_data(as_text=True)


def test_csrf_obligatoire(client):
    r = client.post("/login", data={"login": "admin", "mot_de_passe": "admin"})
    assert r.status_code == 400


def creer_ecole(admin, nom="École Test", **extra):
    d = {"nom": nom, "type": "Primaire", "rne": "9760000A", "nb_classes": 2, "decharge_direction": 0,
         "education_prioritaire": "Hors EP"}
    d.update(extra)
    r = admin.post("/ecoles/nouvelle", data=d)
    assert r.status_code == 302
    return int(r.headers["Location"].rstrip("/").split("/")[-1])


def creer_personnel(admin, ecole_id, nom, **extra):
    d = {"nom": nom, "prenom": "Test", "sexe": "F", "ecole_id": ecole_id, "statut": "PE",
         "dans_organisation": "1", "dispositif": "Solo"}
    d.update(extra)
    r = admin.post("/enseignants/nouveau", data=d)
    assert r.status_code == 302, r.get_data(as_text=True)
    return int(r.headers["Location"].rstrip("/").split("/")[-1])


def test_cycle_complet_ecole_et_stats(admin):
    eid = creer_ecole(admin)
    creer_personnel(admin, eid, "DUPONT", eff_cp="15", eff_ce1="0", dispositif="Duo", salle="1")
    creer_personnel(admin, eid, "DURAND", eff_cp="14", dispositif="Duo", salle="1", sexe="H")
    creer_personnel(admin, eid, "MARTIN", eff_cm1="12", eff_cm2="13", anciennete="7")
    creer_personnel(admin, eid, "REMPLA", statut="TR-ZIL", dans_organisation="")

    page = admin.get(f"/ecoles/{eid}").get_data(as_text=True)
    assert "DUPONT" in page and "REMPLA" in page
    # Effectif total hors ULIS = 15 + 14 + 25 = 54 ; 3 classes en tout ; contrôle classes déclarées (2) ≠ 3
    assert ">54<" in page
    assert "⚠ Écart" in page

    tableau = admin.get("/effectifs/").get_data(as_text=True)
    assert "École Test" in tableau and ">54<" in tableau

    accueil = admin.get("/").get_data(as_text=True)
    assert "École Test" in accueil


def test_grille_effectifs(admin):
    eid = creer_ecole(admin)
    pid = creer_personnel(admin, eid, "GRILLE", eff_cp="10")
    r = admin.post(f"/enseignants/effectifs/{eid}", data={f"eff_cp_{pid}": "12", f"eff_ce1_{pid}": "3",
                                                          f"dispositif_{pid}": "Duo", f"salle_{pid}": "B2"})
    assert r.status_code == 302
    fiche = admin.get(f"/enseignants/{pid}").get_data(as_text=True)
    assert "CP : 12" in fiche and "CE1 : 3" in fiche and "15 élèves" in fiche


def test_suivi_enseignant_et_echeance(admin):
    eid = creer_ecole(admin)
    pid = creer_personnel(admin, eid, "STAGIAIRE", statut="PES1")
    accueil = admin.get("/").get_data(as_text=True)
    assert "STAGIAIRE" in accueil  # à accompagner, sans visite
    r = admin.post(f"/enseignants/{pid}/suivis/nouveau", data={
        "date": "2026-09-15", "type": "Visite conseil", "objet": "Séance de lecture", "statut": "Réalisé",
        "compte_rendu": "Bonne gestion de classe.", "echeance": ""})
    assert r.status_code == 302
    fiche = admin.get(f"/enseignants/{pid}").get_data(as_text=True)
    assert "Séance de lecture" in fiche and "15/09/2026" in fiche
    journal = admin.get("/suivis/").get_data(as_text=True)
    assert "Séance de lecture" in journal
    csv = admin.get("/exports/suivis.csv").get_data(as_text=True)
    assert "Séance de lecture" in csv and csv.startswith("﻿")


def test_eleve_et_suivi(admin):
    eid = creer_ecole(admin)
    pid = creer_personnel(admin, eid, "MAITRE", eff_ce2="20")
    r = admin.post("/eleves/nouveau", data={"nom": "eleve", "prenom": "Un", "ecole_id": eid, "enseignant_id": pid,
                                            "niveau": "CE2", "ppre": "1", "aesh": "Non", "mdph": "Dossier en cours"})
    assert r.status_code == 302
    elid = int(r.headers["Location"].rstrip("/").split("/")[-1])
    r = admin.post(f"/eleves/{elid}/suivis/nouveau", data={
        "date": "2026-10-01", "type": "Équipe éducative", "objet": "Point de rentrée", "statut": "À suivre",
        "echeance": "2026-10-20"})
    assert r.status_code == 302
    liste = admin.get("/eleves/?dispositif=ppre").get_data(as_text=True)
    assert "ELEVE" in liste
    accueil = admin.get("/").get_data(as_text=True)
    assert "Point de rentrée" in accueil  # échéance et derniers suivis


def test_eleve_enseignant_autre_ecole_refuse(admin):
    e1 = creer_ecole(admin, "E1")
    e2 = creer_ecole(admin, "E2", rne="9760001B")
    pid = creer_personnel(admin, e1, "AILLEURS")
    r = admin.post("/eleves/nouveau", data={"nom": "X", "ecole_id": e2, "enseignant_id": pid, "aesh": "Non",
                                            "mdph": "Aucun dossier"})
    assert r.status_code == 200 and "exerce pas dans" in r.get_data(as_text=True)


def test_roles_lecture_et_saisie(admin, client):
    admin.post("/utilisateurs/nouveau", data={"login": "lecteur", "nom": "Lecteur", "role": "lecture",
                                               "mot_de_passe": "motdepasse", "actif": "1"})
    admin.post("/utilisateurs/nouveau", data={"login": "cpc", "nom": "CPC", "role": "saisie",
                                               "mot_de_passe": "motdepasse", "actif": "1"})
    client.post("/logout", data={"_csrf": admin._csrf()})
    lecteur = Session(client)
    assert lecteur.login("lecteur", "motdepasse").status_code == 302
    assert lecteur.get("/ecoles/").status_code == 200
    assert lecteur.get("/ecoles/nouvelle").status_code == 403
    assert lecteur.get("/utilisateurs/").status_code == 403
    client.post("/logout", data={"_csrf": lecteur._csrf()})
    cpc = Session(client)
    assert cpc.login("cpc", "motdepasse").status_code == 302
    assert cpc.get("/ecoles/nouvelle").status_code == 200
    assert cpc.get("/utilisateurs/").status_code == 403
    assert cpc.get("/import/").status_code == 403


def test_admin_ne_peut_pas_se_retirer_ses_droits(admin):
    r = admin.post("/utilisateurs/1/modifier", data={"login": "admin", "nom": "Admin", "role": "saisie", "actif": "1"})
    assert "propres droits" in r.get_data(as_text=True)


def test_separer_nom_prenom():
    assert separer_nom_prenom("BOINALI Habachia") == ("BOINALI", "Habachia")
    assert separer_nom_prenom("Mme SAID ASSOUMANI Rafiktoum") == ("SAID ASSOUMANI", "Rafiktoum")
    assert separer_nom_prenom("SAANDANI Moina (C7)") == ("SAANDANI", "Moina")
    assert separer_nom_prenom("M’CHINDRA Echat-Saïd") == ("M’CHINDRA", "Echat-Saïd")
    assert separer_nom_prenom("VACANT") == ("VACANT", "")


def classeur_fiche_ecole():
    """Construit un classeur minimal au format des fiches école (noms fictifs)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "SYNTHESE"
    ws["A1"] = "TABLEAU DE BORD"
    ws = wb.create_sheet("DEMO PRIM")
    ws["A1"] = "ACADÉMIE DE TEST  —  CIRCONSCRIPTION DE DEMO"
    ws["A2"] = "ECOLE DEMO"
    ws["K2"], ws["M2"] = "Année scolaire :", "2026/2027"
    ws["A3"], ws["E3"], ws["I3"], ws["L3"], ws["M3"] = "Type d'école :", "MATERNELLE", "ÉLÉMENTAIRE", "X", "PRIMAIRE"
    ws["A5"] = "  📋  FICHE SIGNALÉTIQUE"
    ws["A6"], ws["B6"] = "RNE :", "9760099Z"
    ws["A7"], ws["B7"] = "Adresse géographique :", "1 rue du Test"
    ws["A8"], ws["B8"] = "Téléphone / Courriel :", "02 69 00 00 00 / ce.test@ac-test.fr"
    ws["A11"], ws["B11"] = "Directeur :", "Mme DIRECTRICE Test"
    ws["A17"], ws["B17"], ws["H17"], ws["K17"] = "AESH 1 :", "ACCOMP Une", "Tél. AESH :", "0600000000"
    ws["A21"], ws["B21"], ws["H21"], ws["K21"] = "Nombre de salles :", 4, "Nombre total de classes :", 3
    ws["A22"], ws["C22"], ws["J22"], ws["L22"] = "% décharge de direction :", 0.25, "Date de mise à jour :", "2026-09-01"
    entetes = ["Enseignant(e)", "Statut", "Sexe", "Dispositif", "Salle", "PS", "MS", "GS", "CP", "CE1", "CE2",
               "CM1", "CM2", "ULIS", "TOTAL*"]
    for i, h in enumerate(entetes, 1):
        ws.cell(row=25, column=i, value=h)
    lignes = [("ALPHA Une", "PE", "F", "Solo", 1, {"PS": 24}),
              ("BETA Deux", "C", "H", "Duo", 2, {"CP": 15}),
              ("GAMMA Trois", "PES2", "F", "Duo", 2, {"CP": 14, "CE1": 2})]
    for r, (nom, st, sx, disp, salle, eff) in enumerate(lignes, 26):
        ws.cell(row=r, column=1, value=nom); ws.cell(row=r, column=2, value=st); ws.cell(row=r, column=3, value=sx)
        ws.cell(row=r, column=4, value=disp); ws.cell(row=r, column=5, value=salle)
        for niv, n in eff.items():
            ws.cell(row=r, column=entetes.index(niv) + 1, value=n)
    ws["A30"] = "EFFECTIFS TOTAUX PAR NIVEAU  —  Source : ONDE / Arena"
    ws["A32"] = "  👨‍🏫  ENSEIGNANTS NON COMPTABILISÉS DANS L'ORGANISATION PÉDAGOGIQUE  (TR-ZIL / Brigadiers GS)"
    for i, h in enumerate(["N°", "Nom et Prénom", "", "", "", "", "Sexe", "", "Anc. (ans)", "", "Fonction", "", "Observations"], 1):
        if h:
            ws.cell(row=33, column=i, value=h)
    ws["A34"], ws["B34"], ws["G34"], ws["I34"], ws["K34"] = 1, "DELTA Quatre", "F", 3, "TR-ZIL"
    ws["A36"] = "  📊  STATISTIQUES — Calculées automatiquement"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_import_excel_puis_export(admin):
    r = admin.post("/import/", data={"classeur": (classeur_fiche_ecole(), "fiches.xlsx")},
                   content_type="multipart/form-data")
    page = r.get_data(as_text=True)
    assert r.status_code == 200 and "Import terminé" in page
    assert "ECOLE DEMO" in page
    fiche = admin.get("/ecoles/1").get_data(as_text=True)
    assert "9760099Z" in fiche and "ALPHA" in fiche and "DELTA" in fiche and "ACCOMP" in fiche
    assert "25 %" in fiche  # décharge 0.25 -> 25
    assert ">55<" in fiche  # 24 + 15 + 16
    assert "Circonscription De Demo" in fiche

    # Ré-import : pas de doublons
    r = admin.post("/import/", data={"classeur": (classeur_fiche_ecole(), "fiches.xlsx")},
                   content_type="multipart/form-data")
    assert "Écoles mises à jour" in r.get_data(as_text=True)
    liste = admin.get("/enseignants/?q=ALPHA").get_data(as_text=True)
    assert liste.count("ALPHA</strong>") == 1

    # Export du classeur complet
    r = admin.get("/exports/classeur.xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames[:2] == ["BORD", "EFFECTIFS PAR NIVEAUX"] and len(wb.sheetnames) == 3
    ws = wb["EFFECTIFS PAR NIVEAUX"]
    assert ws["A4"].value == "ECOLE DEMO" and ws["O4"].value == 55
    ws_ecole = wb[wb.sheetnames[2]]
    valeurs = [c.value for row in ws_ecole.iter_rows() for c in row if c.value]
    assert "ALPHA Une" in valeurs and "DELTA Quatre" in valeurs
